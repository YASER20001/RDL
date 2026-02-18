"""
KBR RDL Data Harmonizer — Excel Builder
Professional styled Excel exports using openpyxl (no UI framework dependency).
"""

import io
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import openpyxl

from .engine import FILE_TYPES, classify_match, lookup_cfihos, build_export_df

# ── Style constants — KBR corporate navy blue ──
KBR_PRIMARY = "003087"
KBR_DARK = "001D54"
KBR_ACCENT = "00A3E0"
NAVY = "1B3A5C"
WHITE = "FFFFFF"
LIGHT_GRAY = "F8FAFC"
BORDER_GRAY = "E2E8F0"
GREEN_BG = "D1FAE5"
GREEN_FG = "065F46"
RED_BG = "FEE2E2"
RED_FG = "991B1B"
BLUE_BG = "DBEAFE"
BLUE_FG = "1E40AF"
ORANGE_BG = "FFEDD5"
ORANGE_FG = "9A3412"

THIN_BORDER = Border(
    left=Side(style="thin", color=BORDER_GRAY),
    right=Side(style="thin", color=BORDER_GRAY),
    top=Side(style="thin", color=BORDER_GRAY),
    bottom=Side(style="thin", color=BORDER_GRAY),
)


def _style_title_rows(ws, title_text, subtitle_text, num_cols):
    """Add KBR-AMCDE logo + branded title + subtitle rows (rows 1-3)."""
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=num_cols)
    logo_cell = ws.cell(row=1, column=1, value="  KBR-AMCDE  |  RDL Data Harmonizer v3.0")
    logo_cell.font = Font(name="Calibri", size=14, bold=True, color=KBR_ACCENT)
    logo_cell.fill = PatternFill("solid", fgColor=KBR_DARK)
    logo_cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 30

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=num_cols)
    t = ws.cell(row=2, column=1, value=f"  {title_text}")
    t.font = Font(name="Calibri", size=16, bold=True, color=WHITE)
    t.fill = PatternFill("solid", fgColor=KBR_PRIMARY)
    t.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 36

    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=num_cols)
    s = ws.cell(row=3, column=1, value=f"  {subtitle_text}")
    s.font = Font(name="Calibri", size=10, color=WHITE)
    s.fill = PatternFill("solid", fgColor=KBR_DARK)
    s.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[3].height = 24


def _write_headers(ws, row, headers):
    header_font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    header_fill = PatternFill("solid", fgColor=NAVY)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for ci, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=ci, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = THIN_BORDER
    ws.row_dimensions[row].height = 28


def build_excel_bytes(classes, masters, threshold):
    """Build the main harmonization export with professional styling."""
    data_font = Font(name="Calibri", size=10, color=NAVY)
    data_align = Alignment(horizontal="left", vertical="center")
    center_align = Alignment(horizontal="center", vertical="center")
    alt_fill = PatternFill("solid", fgColor=LIGHT_GRAY)
    white_fill = PatternFill("solid", fgColor=WHITE)
    gap_fill = PatternFill("solid", fgColor=RED_BG)
    gap_font = Font(name="Calibri", size=10, bold=True, color=RED_FG)
    nogap_fill = PatternFill("solid", fgColor=GREEN_BG)
    nogap_font = Font(name="Calibri", size=10, bold=True, color=GREEN_FG)

    df = build_export_df(classes, masters)
    total = len(classes)
    non_master_sources = set()
    for c in classes:
        non_master_sources.update(c["matches"].keys())
        non_master_sources.update(c["gaps"])
    non_master_sources = sorted(non_master_sources)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # Sheet 1: Harmonization Results
    ws1 = wb.create_sheet("Harmonization Results")
    cols = list(df.columns)
    master_label = " + ".join(FILE_TYPES.get(m, m) for m in masters)
    compared = ", ".join(FILE_TYPES.get(s, s) for s in non_master_sources)
    _style_title_rows(ws1, "KBR RDL — Harmonization Results",
                      f"Master: {master_label}  |  Compared: {compared}  |  Threshold: {threshold}%  |  Total: {total} classes",
                      len(cols))
    _write_headers(ws1, 4, cols)
    ws1.auto_filter.ref = f"A4:{get_column_letter(len(cols))}4"
    ws1.freeze_panes = "A5"

    for ri, (_, row) in enumerate(df.iterrows(), 5):
        is_alt = (ri % 2 == 0)
        for ci, col_name in enumerate(cols, 1):
            val = row[col_name]
            cell = ws1.cell(row=ri, column=ci, value=val)
            cell.font = data_font
            cell.alignment = center_align if col_name in ("#", "Gap Count", "Masters Cross-Match%") or "Match%" in str(col_name) else data_align
            cell.border = THIN_BORDER
            if col_name == "Gap In":
                if val == "No Gaps":
                    cell.fill = nogap_fill
                    cell.font = nogap_font
                elif val:
                    cell.fill = gap_fill
                    cell.font = gap_font
                else:
                    cell.fill = alt_fill if is_alt else white_fill
            elif "Status" in str(col_name):
                if val == "GAP":
                    cell.fill = gap_fill
                    cell.font = gap_font
                elif val and "Exact" in str(val):
                    cell.fill = nogap_fill
                    cell.font = nogap_font
                else:
                    cell.fill = alt_fill if is_alt else white_fill
            else:
                cell.fill = alt_fill if is_alt else white_fill

    for ci, col_name in enumerate(cols, 1):
        max_len = max(len(str(col_name)), max((len(str(row[col_name] or "")) for _, row in df.iterrows()), default=5))
        ws1.column_dimensions[get_column_letter(ci)].width = min(max_len + 4, 45)

    # Sheet 2: Gap Analysis
    ws2 = wb.create_sheet("Gap Analysis")
    gap_rows = []
    for c in classes:
        for g in c["gaps"]:
            present_in = [FILE_TYPES.get(m, m) for m in masters if m in c["master_entries"]]
            present_in += [FILE_TYPES.get(s, s) for s in c["matches"]]
            gap_rows.append({
                "#": c["index"],
                "Equipment Class": c["canonical_name"],
                "Found In": ", ".join(present_in),
                "Gap In": FILE_TYPES.get(g, g),
                "Action": f"Add '{c['canonical_name']}' to {FILE_TYPES.get(g, g)} or confirm exclusion",
            })
    gap_cols = ["#", "Equipment Class", "Found In", "Gap In", "Action"]
    no_gaps_cnt = sum(1 for c in classes if not c["gaps"])
    _style_title_rows(ws2, "KBR RDL — Gap Analysis",
                      f"{len(gap_rows)} gaps identified  |  {no_gaps_cnt}/{total} classes fully matched",
                      len(gap_cols))
    _write_headers(ws2, 4, gap_cols)
    ws2.auto_filter.ref = f"A4:{get_column_letter(len(gap_cols))}4"
    ws2.freeze_panes = "A5"

    for ri, grow in enumerate(gap_rows, 5):
        is_alt = (ri % 2 == 0)
        for ci, col_name in enumerate(gap_cols, 1):
            val = grow[col_name]
            cell = ws2.cell(row=ri, column=ci, value=val)
            cell.font = data_font
            cell.alignment = center_align if col_name == "#" else data_align
            cell.border = THIN_BORDER
            if col_name == "Gap In":
                cell.fill = gap_fill
                cell.font = gap_font
            else:
                cell.fill = alt_fill if is_alt else white_fill
    gap_widths = [6, 40, 40, 25, 55]
    for ci, w in enumerate(gap_widths, 1):
        ws2.column_dimensions[get_column_letter(ci)].width = w

    # Sheet 3: Summary
    ws3 = wb.create_sheet("Summary")
    _style_title_rows(ws3, "KBR RDL — Summary Report", f"Generated from harmonization of {total} equipment classes", 2)
    _write_headers(ws3, 4, ["Metric", "Value"])

    no_gaps = sum(1 for c in classes if not c["gaps"])
    summary_data = [
        ("Total Equipment Classes", total),
        ("Master Reference", master_label),
        ("Compared Against", compared),
        ("Match Threshold", f"{threshold}%"),
        ("Classes with No Gaps", f"{no_gaps} / {total}"),
        ("Total Gap Entries", sum(len(c["gaps"]) for c in classes)),
    ]
    for src in non_master_sources:
        cnt = sum(1 for c in classes if src in c["matches"])
        pct = int(round(cnt / total * 100)) if total else 0
        gap_cnt = total - cnt
        summary_data.append((FILE_TYPES.get(src, src), f"Matched: {cnt}/{total} ({pct}%) — Gaps: {gap_cnt}"))

    for ri, (metric, value) in enumerate(summary_data, 5):
        is_alt = (ri % 2 == 0)
        mc = ws3.cell(row=ri, column=1, value=metric)
        mc.font = Font(name="Calibri", size=10, bold=True, color=NAVY)
        mc.fill = alt_fill if is_alt else white_fill
        mc.border = THIN_BORDER
        vc = ws3.cell(row=ri, column=2, value=value)
        vc.font = Font(name="Calibri", size=10, color=NAVY)
        vc.fill = alt_fill if is_alt else white_fill
        vc.border = THIN_BORDER
    ws3.column_dimensions["A"].width = 30
    ws3.column_dimensions["B"].width = 50

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def build_enriched_master_excel(selected_additions, masters, files, threshold, aramco_attrs=None):
    """Build enriched master Excel with original + suggested additions + class attributes."""
    title_font = Font(name="Calibri", size=16, bold=True, color=WHITE)
    title_fill = PatternFill("solid", fgColor=KBR_PRIMARY)
    sub_fill = PatternFill("solid", fgColor=KBR_DARK)
    header_font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    header_fill = PatternFill("solid", fgColor=NAVY)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    data_font = Font(name="Calibri", size=10, color=NAVY)
    data_align = Alignment(horizontal="left", vertical="center")
    center_align = Alignment(horizontal="center", vertical="center")
    orig_fill = PatternFill("solid", fgColor=WHITE)
    sugg_fill = PatternFill("solid", fgColor=BLUE_BG)
    alt_fill = PatternFill("solid", fgColor=LIGHT_GRAY)
    exact_fill = PatternFill("solid", fgColor=GREEN_BG)
    exact_font = Font(name="Calibri", size=10, bold=True, color=GREEN_FG)
    partial_fill = PatternFill("solid", fgColor=ORANGE_BG)
    partial_font = Font(name="Calibri", size=10, bold=True, color=ORANGE_FG)
    status_orig_fill = PatternFill("solid", fgColor=GREEN_BG)
    status_orig_font = Font(name="Calibri", size=10, bold=True, color=GREEN_FG)
    status_sugg_fill = PatternFill("solid", fgColor=BLUE_BG)
    status_sugg_font = Font(name="Calibri", size=10, bold=True, color=BLUE_FG)

    cfihos_records = files["cfihos"]["records"] if "cfihos" in files else []

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    columns = ["#", "Status", "Id", "Name", "CFIHOS Code", "CFIHOS Name", "CFIHOS Match"]
    col_widths = [6, 28, 18, 40, 16, 40, 14]

    for m in masters:
        if m not in files:
            continue
        original = files[m]["records"]
        label = FILE_TYPES.get(m, m)
        ws = wb.create_sheet(title=f"{label} Enriched")
        nc = len(columns)

        # Logo row
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=nc)
        lc = ws.cell(row=1, column=1, value="  KBR-AMCDE  |  RDL Data Harmonizer v3.0")
        lc.font = Font(name="Calibri", size=14, bold=True, color=KBR_ACCENT)
        lc.fill = PatternFill("solid", fgColor=KBR_DARK)
        lc.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[1].height = 30

        # Title
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=nc)
        tc = ws.cell(row=2, column=1, value=f"  {label} — Enriched Master")
        tc.font = title_font
        tc.fill = title_fill
        tc.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[2].height = 36

        # Subtitle
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=nc)
        sc = ws.cell(row=3, column=1,
                     value=f"  Original: {len(original)} classes  |  Suggested Additions: {len(selected_additions)}  |  New Total: {len(original) + len(selected_additions)}  |  Threshold: {threshold}%")
        sc.font = Font(name="Calibri", size=10, color=WHITE)
        sc.fill = sub_fill
        sc.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[3].height = 24

        # Headers
        for ci, col_name in enumerate(columns, 1):
            cell = ws.cell(row=4, column=ci, value=col_name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = THIN_BORDER
        ws.row_dimensions[4].height = 28
        ws.auto_filter.ref = f"A4:{get_column_letter(nc)}4"

        row_num = 5
        counter = 1

        # Original records
        for r in original:
            cfihos_code = r.get("cfihos_ref", "")
            cfihos_name = ""
            cfihos_match_label = ""
            if cfihos_code and cfihos_code != "nan" and cfihos_records:
                for cr in cfihos_records:
                    if cr["id"] == cfihos_code:
                        cfihos_name = cr["name"]
                        cfihos_match_label = "Exact"
                        break
                if not cfihos_name:
                    cfihos_match_label = "Ref Only"
            if not cfihos_code or cfihos_code == "nan":
                cfihos_code, cfihos_name, cfihos_match_label = lookup_cfihos(r["name"], cfihos_records, threshold)

            values = [counter, "Original", r["id"], r["name"],
                      cfihos_code if cfihos_code != "nan" else "", cfihos_name, cfihos_match_label]
            is_alt = (counter % 2 == 0)
            for ci, val in enumerate(values, 1):
                cell = ws.cell(row=row_num, column=ci, value=val)
                cell.font = data_font
                cell.alignment = center_align if ci in (1, 7) else data_align
                cell.border = THIN_BORDER
                if ci == 2:
                    cell.fill = status_orig_fill
                    cell.font = status_orig_font
                    cell.alignment = center_align
                elif ci == 7 and val:
                    if val == "Exact":
                        cell.fill = exact_fill
                        cell.font = exact_font
                    else:
                        cell.fill = partial_fill
                        cell.font = partial_font
                else:
                    cell.fill = alt_fill if is_alt else orig_fill
            counter += 1
            row_num += 1

        # Separator
        ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=len(columns))
        sep = ws.cell(row=row_num, column=1, value="  SUGGESTED ADDITIONS")
        sep.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
        sep.fill = PatternFill("solid", fgColor="2563EB")
        sep.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[row_num].height = 26
        row_num += 1

        # Suggested additions
        for rec in selected_additions:
            cfihos_code, cfihos_name, cfihos_match_label = lookup_cfihos(rec["name"], cfihos_records, threshold)
            src_label = FILE_TYPES.get(rec["source"], rec["source"])
            values = [counter, f"Suggested from {src_label}", f"NEW-{rec['id']}", rec["name"],
                      cfihos_code, cfihos_name, cfihos_match_label]
            is_alt = (counter % 2 == 0)
            for ci, val in enumerate(values, 1):
                cell = ws.cell(row=row_num, column=ci, value=val)
                cell.font = data_font
                cell.alignment = center_align if ci in (1, 7) else data_align
                cell.border = THIN_BORDER
                if ci == 2:
                    cell.fill = status_sugg_fill
                    cell.font = status_sugg_font
                    cell.alignment = center_align
                elif ci == 7 and val:
                    if val == "Exact":
                        cell.fill = exact_fill
                        cell.font = exact_font
                    else:
                        cell.fill = partial_fill
                        cell.font = partial_font
                else:
                    cell.fill = sugg_fill if not is_alt else PatternFill("solid", fgColor="C7D2FE")
            counter += 1
            row_num += 1

        for ci, w in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(ci)].width = w
        ws.freeze_panes = "A5"

    # Summary sheet
    ws_sum = wb.create_sheet(title="Enrichment Summary")
    orig_total = sum(len(files[m]["records"]) for m in masters if m in files)

    ws_sum.merge_cells("A1:B1")
    lc = ws_sum.cell(row=1, column=1, value="  KBR-AMCDE  |  RDL Data Harmonizer v3.0")
    lc.font = Font(name="Calibri", size=14, bold=True, color=KBR_ACCENT)
    lc.fill = PatternFill("solid", fgColor=KBR_DARK)
    lc.alignment = Alignment(horizontal="left", vertical="center")
    ws_sum.row_dimensions[1].height = 30

    ws_sum.merge_cells("A2:B2")
    t = ws_sum.cell(row=2, column=1, value="  Enrichment Summary")
    t.font = Font(name="Calibri", size=16, bold=True, color=WHITE)
    t.fill = PatternFill("solid", fgColor=KBR_PRIMARY)
    t.alignment = Alignment(horizontal="left", vertical="center")
    ws_sum.row_dimensions[2].height = 36

    summary_data = [
        ("Original Master", " + ".join(FILE_TYPES.get(m, m) for m in masters)),
        ("Original Classes", orig_total),
        ("Suggested Additions", len(selected_additions)),
        ("New Total", orig_total + len(selected_additions)),
        ("Match Threshold", f"{threshold}%"),
    ]
    for src in sorted(set(r["source"] for r in selected_additions)):
        cnt = sum(1 for r in selected_additions if r["source"] == src)
        summary_data.append((f"Added from {FILE_TYPES.get(src, src)}", cnt))

    _write_headers(ws_sum, 4, ["Metric", "Value"])
    for ri, (metric, value) in enumerate(summary_data, 5):
        is_alt = ri % 2 == 0
        mc = ws_sum.cell(row=ri, column=1, value=metric)
        mc.font = Font(name="Calibri", size=10, bold=True, color=NAVY)
        mc.fill = PatternFill("solid", fgColor=LIGHT_GRAY) if is_alt else PatternFill("solid", fgColor=WHITE)
        mc.border = THIN_BORDER
        vc = ws_sum.cell(row=ri, column=2, value=value)
        vc.font = Font(name="Calibri", size=10, color=NAVY)
        vc.fill = PatternFill("solid", fgColor=LIGHT_GRAY) if is_alt else PatternFill("solid", fgColor=WHITE)
        vc.border = THIN_BORDER
    ws_sum.column_dimensions["A"].width = 30
    ws_sum.column_dimensions["B"].width = 35

    # Class Attributes sheet
    if aramco_attrs is not None and not aramco_attrs.empty and "Class_Id" in aramco_attrs.columns:
        ws_attr = wb.create_sheet(title="Class Attributes")
        attr_cols_list = ["Class_Id", "Name", "Description", "Presence", "Size",
                          "Discipline", "UomClassId", "UomRequire", "ValidationRule", "Group_Id"]
        attr_cols_list = [c for c in attr_cols_list if c in aramco_attrs.columns]
        nc = len(attr_cols_list)

        ws_attr.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(nc, 2))
        lc = ws_attr.cell(row=1, column=1, value="  KBR-AMCDE  |  RDL Data Harmonizer v3.0")
        lc.font = Font(name="Calibri", size=14, bold=True, color=KBR_ACCENT)
        lc.fill = PatternFill("solid", fgColor=KBR_DARK)
        lc.alignment = Alignment(horizontal="left", vertical="center")
        ws_attr.row_dimensions[1].height = 30

        ws_attr.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(nc, 2))
        tc = ws_attr.cell(row=2, column=1, value="  ISM Functional Class Attributes — Organized by Class")
        tc.font = Font(name="Calibri", size=16, bold=True, color=WHITE)
        tc.fill = PatternFill("solid", fgColor=KBR_PRIMARY)
        tc.alignment = Alignment(horizontal="left", vertical="center")
        ws_attr.row_dimensions[2].height = 36

        all_class_ids = {}
        for m_key in masters:
            if m_key in files:
                for r in files[m_key]["records"]:
                    all_class_ids.setdefault(r["name"], set()).add(str(r["id"]).strip())
        for rec in selected_additions:
            all_class_ids.setdefault(rec["name"], set()).add(str(rec["id"]).strip())

        row_num = 3
        for class_name in sorted(all_class_ids.keys()):
            ids = all_class_ids[class_name]
            mask = aramco_attrs["Class_Id"].astype(str).str.strip().isin(ids)
            subset = aramco_attrs[mask]
            if subset.empty:
                continue

            ws_attr.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=nc)
            sep = ws_attr.cell(row=row_num, column=1,
                               value=f"  {class_name}  ({', '.join(sorted(ids))})  —  {len(subset)} attributes")
            sep.font = Font(name="Calibri", size=11, bold=True, color=WHITE)
            sep.fill = PatternFill("solid", fgColor=KBR_PRIMARY)
            sep.alignment = Alignment(horizontal="left", vertical="center")
            ws_attr.row_dimensions[row_num].height = 26
            row_num += 1

            _write_headers(ws_attr, row_num, attr_cols_list)
            row_num += 1

            for _, attr_row in subset.iterrows():
                is_alt = (row_num % 2 == 0)
                for ci, col_name in enumerate(attr_cols_list, 1):
                    val = attr_row.get(col_name, "")
                    if pd.isna(val):
                        val = ""
                    cell = ws_attr.cell(row=row_num, column=ci, value=str(val))
                    cell.font = Font(name="Calibri", size=10, color=NAVY)
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                    cell.border = THIN_BORDER
                    if col_name == "Presence":
                        sv = str(val).strip().lower()
                        if sv in ("mandatory", "required", "m"):
                            cell.fill = PatternFill("solid", fgColor=GREEN_BG)
                            cell.font = Font(name="Calibri", size=10, bold=True, color=GREEN_FG)
                        elif sv in ("optional", "o"):
                            cell.fill = PatternFill("solid", fgColor=BLUE_BG)
                            cell.font = Font(name="Calibri", size=10, color=BLUE_FG)
                        else:
                            cell.fill = PatternFill("solid", fgColor=LIGHT_GRAY) if is_alt else PatternFill("solid", fgColor=WHITE)
                    else:
                        cell.fill = PatternFill("solid", fgColor=LIGHT_GRAY) if is_alt else PatternFill("solid", fgColor=WHITE)
                row_num += 1
            row_num += 1

        attr_widths = {"Class_Id": 14, "Name": 35, "Description": 40, "Presence": 12,
                       "Size": 8, "Discipline": 15, "UomClassId": 14, "UomRequire": 12,
                       "ValidationRule": 20, "Group_Id": 14}
        for ci, col_name in enumerate(attr_cols_list, 1):
            ws_attr.column_dimensions[get_column_letter(ci)].width = attr_widths.get(col_name, 18)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
